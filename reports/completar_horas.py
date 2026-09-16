from pathlib import Path
from datetime import datetime, timedelta
from shutil import copy2
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.comments import Comment

source = Path(r'C:\Users\kevin\Desktop\OctavoSemestre\PRACTICAS\horas_practica.xlsx')
out = Path(__file__).parent / 'bitacoras_agosto_septiembre'
out.mkdir(parents=True, exist_ok=True)
backup = out / 'horas_practica_original.xlsx'
if not backup.exists():
    copy2(source, backup)
wb = load_workbook(backup)
ws = wb.active
original = [[ws.cell(r,c).value for c in range(1,4)] for r in range(2,6)]
activities = [
    'Mejoré el flujo de los informes y la gestión de usuarios. Agregué la tabla de docentes y realicé ajustes en la interfaz.',
    'Avancé en el procesamiento de los archivos Excel y en la organización de la información para generar los informes.',
    'Completé el flujo de generación de informes Excel e integré las hojas y los datos necesarios para el reporte.',
    'Trabajé en la plantilla del informe Word y en la incorporación de los gráficos generados a partir del Excel.',
    'Avancé en la integración de la generación del informe Word con la información procesada en Excel.',
    'Mejoré la presentación visual de los reportes Word y la organización de su contenido.',
    'Ajusté la plantilla del informe, mejoré las gráficas y optimicé la generación de los archivos Word y Excel.',
    'Mejoré la interfaz del generador Word y los iconos de Excel. Ajusté la tabla de contenido y la conversión a PDF para Windows y Mac.',
    'Corregí detalles de la tabla de contenido y de la portada al generar el PDF. Incorporé la mascota Tomy con ayuda contextual en cuatro módulos.',
    'Revisión propuesta del flujo de generación de informes Excel, Word y PDF y de la ayuda contextual; confirmar la actividad realizada este día.',
]
day = datetime(2026,8,25)
row = 6
for activity in activities:
    while day.weekday() >= 5:
        day += timedelta(days=1)
    ws.cell(row,1,day)
    ws.cell(row,2,[2,5,3,4,5][day.weekday()])
    ws.cell(row,3,activity)
    ws.merge_cells(start_row=row,start_column=3,end_row=row,end_column=9)
    ws.cell(row,3).comment = Comment('Distribución propuesta a partir del historial del proyecto y los archivos disponibles. Confirmar actividad y fecha antes de presentar. Las horas corresponden al horario programado.', 'Revisión')
    day += timedelta(days=1)
    row += 1
ws.cell(row,1,'TOTAL')
ws.cell(row,2,f'=SUM(B2:B{row-1})')
ws.merge_cells(start_row=row+2,start_column=1,end_row=row+3,end_column=9)
ws.cell(row+2,1,'PARA REVISIÓN: Se conservaron los cuatro registros originales. Las actividades del 25/08 al 07/09 son una distribución propuesta basada en los avances del proyecto; validar antes de presentar. Total según horario: 52 horas. La actividad del 07/09 requiere confirmación.')
ws.cell(row+2,1).alignment = Alignment(wrap_text=True,vertical='center')
ws.cell(row+2,1).font = Font(size=10,italic=True,color='805000')
ws.row_dimensions[row+2].height = 32
ws.row_dimensions[row+3].height = 24
ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 10
for col in 'CDEFGHI':
    ws.column_dimensions[col].width = 12
edge = Side(style='thin',color='D9E1F2')
for r in range(1,row+1):
    ws.row_dimensions[r].height = 48 if 1 < r < row else 25
    for c in range(1,10):
        cell = ws.cell(r,c)
        cell.font = Font(name='Calibri',size=11,bold=r in [1,row])
        cell.alignment = Alignment(horizontal='left' if c>=3 else 'center',vertical='center',wrap_text=True)
        cell.border = Border(bottom=edge)
        if r in [1,row]:
            cell.fill = PatternFill('solid',fgColor='D9EAF7')
    if 1 < r < row:
        ws.cell(r,1).number_format = 'dd/mm/yyyy'
ws.freeze_panes = 'C2'
ws.print_options.horizontalCentered = True
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = 'landscape'
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.print_area = f'A1:I{row+3}'
target = out / 'horas_practica_completado.xlsx'
wb.save(target)
check = load_workbook(target)
s = check.active
assert [[s.cell(r,c).value for c in range(1,4)] for r in range(2,6)] == original
assert sum(s.cell(r,2).value for r in range(2,16)) == 52
assert s['A15'].value == datetime(2026,9,7)
assert all(s.cell(r,3).value for r in range(2,16))
assert s['B16'].value == '=SUM(B2:B15)'
print(f'Verificado: 14 días, 52 horas, cuatro registros originales conservados. {target}')
