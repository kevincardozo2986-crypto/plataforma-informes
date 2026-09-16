from copy import copy
from datetime import datetime, timedelta
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.comments import Comment

folder = Path(__file__).parent / 'bitacoras_agosto_septiembre'
wb = load_workbook(folder / 'horas_practica_completado.xlsx')
ws = wb.active
preserved = [[ws.cell(r,c).value for c in range(1,4)] for r in range(2,16)]
for merged in list(ws.merged_cells.ranges):
    if merged.min_row >= 16:
        ws.unmerge_cells(str(merged))
ws.delete_rows(16, ws.max_row-15)
total = 52
day = datetime(2026,9,8)
row = 16
while total < 192:
    if day.weekday() < 5:
        hours = min([2,5,3,4,5][day.weekday()],192-total)
        for col in range(1,10):
            ws.cell(row,col)._style = copy(ws.cell(15,col)._style)
        ws.cell(row,1,day)
        ws.cell(row,2,hours)
        ws.cell(row,3,'Pendiente de registrar la actividad realizada.')
        ws.cell(row,3).comment = Comment('Horas programadas según el horario; no constituyen confirmación de asistencia ni de actividad realizada.', 'Planificación')
        ws.merge_cells(start_row=row,start_column=3,end_row=row,end_column=9)
        ws.row_dimensions[row].height = 32
        total += hours
        row += 1
    day += timedelta(days=1)
last = row-1
ws.cell(row,1,'TOTAL')
ws.cell(row,2,f'=SUM(B2:B{last})')
for col in range(1,10):
    ws.cell(row,col)._style = copy(ws.cell(1,col)._style)
ws.merge_cells(start_row=row+2,start_column=1,end_row=row+4,end_column=9)
note = ws.cell(row+2,1,'PLANIFICACIÓN DE 192 HORAS: se aplica el horario de lunes a viernes sin descontar festivos, ausencias ni cambios de jornada. Se conservan los registros anteriores; las actividades del 25/08 al 07/09 siguen siendo propuestas para revisión. Desde el 08/09 las actividades quedan pendientes. El último día se programan 2 de las 3 horas habituales para completar exactamente 192 horas.')
note.alignment = Alignment(wrap_text=True,vertical='center')
note.font = Font(name='Calibri',size=11,italic=True,color='805000')
for r in range(row+2,row+5):
    ws.row_dimensions[r].height = 22
ws.page_setup.fitToHeight = 0
ws.print_title_rows = '1:1'
ws.print_area = f'A1:I{row+4}'
target = folder / 'horas_practica_192_horas.xlsx'
wb.save(target)
check = load_workbook(target).active
assert [[check.cell(r,c).value for c in range(1,4)] for r in range(2,16)] == preserved
assert sum(check.cell(r,2).value for r in range(2,last+1)) == 192
days = [check.cell(r,1).value for r in range(2,last+1)]
assert len(days) == len(set(days))
assert all(d.weekday()<5 for d in days)
for r in range(2,last):
    assert check.cell(r,2).value == [2,5,3,4,5][check.cell(r,1).value.weekday()]
assert check.cell(last,2).value == 2
print(f'Archivo: {target}; días: {len(days)}; total: 192; fin: {days[-1]:%d/%m/%Y}; última jornada: 2 horas.')
