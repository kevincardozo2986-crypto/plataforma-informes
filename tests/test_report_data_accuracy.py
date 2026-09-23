from datetime import datetime
from contextlib import contextmanager

import pandas as pd
import pytest
from openpyxl import load_workbook

from app.services.csv_service import CSVValidationError, iter_csv_chunks
from app.services.excel_service import ExcelProcess


def test_period_filter_handles_boundaries_invalid_dates_and_chunks(tmp_path):
    dates = ["2025-12-31 23:59", "2026-01-01", "2026-06-30 23:59",
             "2026-07-01", "2026-12-31 23:59", "2027-01-01"]
    timestamps = [int(pd.Timestamp(d, tz="America/Bogota").timestamp()) for d in dates]
    path = tmp_path / "events.csv"
    pd.DataFrame({"FechaUnix": ["invalid", *timestamps], "idusuario": range(7)}).to_csv(path, index=False)
    for prepare in (True, False):
        for period, expected in (("2026-1", [2, 3]), ("2026-2", [4, 5])):
            stats = {}
            data = pd.concat(iter_csv_chunks(path, chunksize=1, prepare=prepare, period=period, stats=stats))
            assert data.idusuario.tolist() == expected
            assert stats == {"outside_period": 4, "invalid_dates": 1}


def test_empty_period_does_not_replace_existing_workbook(tmp_path):
    path = tmp_path / "events.csv"
    pd.DataFrame({"FechaUnix": [1704067200]}).to_csv(path, index=False)
    process = ExcelProcess(tmp_path / "report.xlsx")
    process.create_original(pd.DataFrame({"original": ["preservar"]}))
    before = process.path.read_bytes()
    with pytest.raises(CSVValidationError, match="No hay registros"):
        process.create_original_from_chunks(iter_csv_chunks(path, period="2026-1"))
    assert process.path.read_bytes() == before


@pytest.mark.parametrize("reopen", [False, True])
def test_distinct_students_teachers_and_years(tmp_path, reopen):
    data = pd.DataFrame({
        "curso": ["A"] * 6,
        "idusuario": [10, 10, 20, 21, 20, 10],
        "usuario": ["Ana", "Ana", "Juan", "Juan", "Juan corregido", "Ana"],
        "rol": ["student", "student", "editingteacher", "editingteacher", "editingteacher", "student"],
        "Mes": [2, 3, 2, 2, 2, 2], "Dia": [1, 1, 1, 2, 3, 1],
        "Fecha": [datetime(2026, 2, 1), datetime(2026, 3, 1), datetime(2026, 2, 1),
                  datetime(2026, 2, 2), datetime(2026, 2, 3), datetime(2025, 2, 1)],
        "accion": ["viewed"] * 6,
    })
    process = ExcelProcess(tmp_path / "report.xlsx")
    process.create_original_from_chunks([data])
    expected_summary = process._report_summary_cache
    if reopen:
        process = ExcelProcess(process.path)
    process.crear_tabla_docentes()
    with load_workbook_context(process.path) as book:
        teachers = list(book["Tabla Dinamica Docentes"].iter_rows(min_row=2, max_row=3, values_only=True))
        assert teachers == [("A", "Juan corregido", 2, 2), ("A", "Juan", 1, 1)]
        summary = process._calcular_resumen_informe(book["Original"])
        assert summary == expected_summary
        assert summary["estudiantes"] == 1
        assert summary["docentes"] == 2
        assert summary["promedio_dias_estudiantes"] == 3
        assert sorted(t["dias"] for t in summary["docentes_destacados"]) == [1, 2]
    process.crear_tabla_estudiantes("Prueba", "2026-1")
    with load_workbook_context(process.path) as book:
        row = list(book["Tabla Dinamica Estudiantes"].iter_rows(min_row=4, max_row=4, values_only=True))[0]
        assert row[:7] == ("A", 2, 1, 3, 1, 1, 1)


def test_nullable_fields_do_not_crash_and_missing_ids_are_not_people(tmp_path):
    frame = pd.DataFrame({
        "curso": ["A", "A", pd.NA], "usuario": ["Ana", pd.NA, pd.NA],
        "idusuario": [10, pd.NA, pd.NA], "rol": ["student", "student", pd.NA],
        "Mes": [2, 2, pd.NA], "Dia": [1, 2, pd.NA],
        "accion": ["viewed", "created", pd.NA],
    }).convert_dtypes()
    process = ExcelProcess(tmp_path / "report.xlsx")
    process.create_original_from_chunks([frame])
    assert process.missing_user_ids == 1
    assert process._report_summary_cache["usuarios_unicos"] == 1
    assert process._report_summary_cache["total_eventos"] == 2
    assert process._activity_summary_cache["A"]["created"] == 1
    with load_workbook_context(process.path) as book:
        assert process._calcular_resumen_informe(book["Original"]) == process._report_summary_cache


@contextmanager
def load_workbook_context(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        yield book
    finally:
        book.close()
